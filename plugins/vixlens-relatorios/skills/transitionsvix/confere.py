"""Confere o relatorio Transitions do mes: compacto (FABRICIO) fecha com o completo, e compara com o mes anterior.
Uso: python .claude/skills/transitionsvix/confere.py 2026-08
Le entregaveis/relatorios/Transitions-<Mes>-<Ano>.xlsx e "... - FABRICIO.xlsx". Nao toca no banco.
"""
import sys, os, glob
from datetime import date
import openpyxl

MESES = ['janeiro','fevereiro','marco','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']
MESES_ACENTO = {'marco': 'Março'}

def nome_mes(ym):
    y, m = ym.split('-'); n = MESES[int(m)-1]
    cap = MESES_ACENTO.get(n, n.capitalize())
    return f"{cap}-{y}"

def arquivos(ym):
    base = os.path.join('entregaveis', 'relatorios')
    full = os.path.join(base, f"Transitions-{nome_mes(ym)}.xlsx")
    comp = os.path.join(base, f"Transitions-{nome_mes(ym)} - FABRICIO.xlsx")
    if not os.path.exists(full):  # tolera variacao de acento no nome
        c = glob.glob(os.path.join(base, f"Transitions-*-{ym[:4]}.xlsx"))
        full = next((f for f in c if MESES[int(ym[5:])-1][:3] in os.path.basename(f).lower()), full)
    return full, comp

def le_completo(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    cli = [r for r in wb['Clientes Transitions'].iter_rows(min_row=5, values_only=True) if isinstance(r[0], (int, float))]
    sem = [r for r in wb['NAO levaram Transitions'].iter_rows(min_row=6, values_only=True) if isinstance(r[0], (int, float))]
    det = [r for r in wb['Cliente x modelo'].iter_rows(min_row=5, values_only=True) if isinstance(r[0], (int, float))]
    mod = [r for r in wb['Modelos Transitions'].iter_rows(min_row=5, values_only=True) if isinstance(r[0], (int, float))]
    return dict(cli=cli, sem=sem, det=det, mod=mod,
                trans=sum(r[7] for r in cli), lentes_cli=sum(r[9] for r in cli), lentes_sem=sum(r[9] for r in sem))

def le_compacto(path):
    """O compacto do Fabricio tem UMA aba (CEO 21/09): todos os clientes, Transitions DESC,
    quem nao levou no fim com 0. Antes eram 3 abas e a pessoa abria na errada."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    todas = [r for r in ws.iter_rows(min_row=2, values_only=True) if isinstance(r[0], (int, float))]
    com = [r for r in todas if r[5]]
    sem = [r for r in todas if not r[5]]
    return dict(nomes=wb.sheetnames, ativa=wb.active.title, todas=todas, fc=com, fn=sem)


def main():
    ym = sys.argv[1] if len(sys.argv) > 1 else None
    if not ym:
        t = date.today(); m = t.month - 1 or 12; y = t.year if t.month > 1 else t.year - 1
        ym = f"{y}-{m:02d}"
    full, comp = arquivos(ym)
    for p in (full, comp):
        if not os.path.exists(p):
            print(f"FALTA: {p}"); sys.exit(2)
    A = le_completo(full); B = le_compacto(comp)
    ok = True
    def chk(cond, msg):
        nonlocal ok
        print(("OK   " if cond else "ERRO ") + msg); ok = ok and cond
    chk(len(B['nomes']) == 1, f"1 aba so, sem aba errada para abrir: {B['nomes']}")
    chk(len(B['todas']) == len(A['cli']) + len(A['sem']),
        f"todos os clientes na aba ({len(B['todas'])} = {len(A['cli'])} com + {len(A['sem'])} sem)")
    chk([r[5] for r in B['todas']] == sorted((r[5] for r in B['todas']), reverse=True),
        "ordenado por Lentes Transitions, maior -> menor (zeros no fim)")
    chk([r[1] for r in B['fc']] == [r[1] for r in A['cli']], f"quem levou = Clientes Transitions ({len(B['fc'])} clientes, mesma ordem)")
    chk(sum(r[5] for r in B['fc']) == A['trans'], f"lentes Transitions compacto = completo ({A['trans']})")
    chk(sum(r[6] for r in B['fc']) == A['lentes_cli'], f"lentes totais dos compradores = completo ({A['lentes_cli']})")
    chk([r[1] for r in B['fn']] == [r[1] for r in A['sem']], f"quem nao levou = NAO levaram ({len(B['fn'])} clientes, mesma ordem)")
    chk(sum(r[6] for r in B['fn']) == A['lentes_sem'], f"lentes dos que nao levaram = completo ({A['lentes_sem']})")
    chk(len(set(r[1] for r in B['todas'])) == len(B['todas']), "nenhum cliente repetido")
    chk(all(abs(r[7] - (r[5] / r[6] if r[6] else 0)) < 1e-9 for r in B['todas']), "% Transitions = trans/lentes")
    chk(sum(r[6] for r in A['det']) == A['trans'], f"Cliente x modelo soma {sum(r[6] for r in A['det'])} = total {A['trans']}")
    print()
    print(f"{ym}: {A['trans']:,} lentes Transitions ({A['trans']/2:,.1f} pares) | {len(A['mod'])} modelos | "
          f"{len(A['cli'])} compraram ({A['lentes_cli']:,} lentes) | {len(A['sem'])} NAO levaram ({A['lentes_sem']:,} lentes)")
    top = A['cli'][:3]; print("  top compradores:", "; ".join(f"{r[2]} ({r[7]})" for r in top))
    top = A['sem'][:3]; print("  maiores sem Transitions:", "; ".join(f"{r[2]} ({r[9]} lentes)" for r in top))
    # mes anterior
    y, m = int(ym[:4]), int(ym[5:]); pm = f"{y if m > 1 else y-1}-{(m-1) or 12:02d}"
    pf, _ = arquivos(pm)
    if os.path.exists(pf):
        P = le_completo(pf)
        print(f"  vs {pm}: {P['trans']:,} lentes | {len(P['cli'])} compraram | {len(P['sem'])} NAO levaram ({P['lentes_sem']:,} lentes)")
    else:
        print(f"  (sem arquivo do mes anterior {pm} para comparar)")
    print("\nRESULTADO:", "FECHA" if ok else "NAO FECHA - nao entregar")
    sys.exit(0 if ok else 1)

if __name__ == '__main__':
    main()
