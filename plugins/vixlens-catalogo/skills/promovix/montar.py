"""Lê o Excel da Promovix, confere os dados e gera os scripts use_figma.

Uso:
    python montar.py Promovix.xlsx --saida pasta/

Grava pasta/lote_NN.js (construtor + páginas, cada lote abaixo do limite do
use_figma) e pasta/resumo.json. Sai com código 1 se houver ERRO; AVISO não bloqueia.

Combo (opcional): aba "Combos" + chaves combo_* na aba Mês + bloco COMBO na aba Páginas.
Sem isso, a Promovix sai exatamente como antes (sem página de combo).
"""
import argparse, json, re, sys
from pathlib import Path
from openpyxl import load_workbook

AQUI = Path(__file__).resolve().parent
GEN_S = ['Cinza', 'Marrom', 'Verde', 'Ametista', 'Safira', 'Âmbar', 'Esmeralda', 'Rubi']
COLORS = ['Marrom', 'G15', 'Black']
ESPELHADO = ['Prata', 'Dourado', 'Azul', 'Rosa']
CONHECIDAS = set(GEN_S + COLORS + ESPELHADO)
ESPECIAIS = {'CAPA', 'FAIXA_50', 'FAIXA_MONTAGEM', 'VIXCLUB', 'COMBO'}
MARCAS_COMBO = {'Kodak', 'Optview', 'Essilor'}
PAINEIS = {'VIXCLUB', 'CAMPANHA_FREEVIX', 'NENHUM'}
QR_URL_PADRAO = 'https://vixlens.com.br/vix-club#segundo-par'
LIMITE, TOPO, TOPO_CAPA = 796, 20, 102

erros, avisos = [], []

def txt(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def sn(v): return txt(v).upper() in ('S', 'SIM', 'X', '1', 'TRUE')

def num(v):
    if v is None or txt(v) == '': return None
    if isinstance(v, (int, float)): return float(v)
    s = txt(v).replace('R$', '').replace(' ', '')
    if ',' in s: s = s.replace('.', '').replace(',', '.')
    return float(s)

def dec(v):
    """+6.00 / -10.25 / 0.00 — sempre duas casas, ponto, sinal explícito."""
    if abs(v) < 1e-9: return '0.00'
    return ('+' if v > 0 else '-') + f'{abs(v):.2f}'

def reais(v):
    s = f'{v:,.2f}'
    return s.replace(',', '#').replace('.', ',').replace('#', '.')

# Títulos amigáveis do Excel (01/10/2026) -> nomes internos. Os nomes antigos continuam valendo.
ALIAS = {
    'Família': 'familia', 'Subtítulo': 'tipo', 'Cor (hex)': 'cor', 'Pílulas do cabeçalho': 'pilulas',
    'Colunas de preço': 'colunas_preco', 'Colunas em promoção': 'colunas_promo', 'Regra 50%': 'regra_50',
    'Faixa ÍNDICE': 'separador_indice', 'Coluna Tratamento': 'coluna_tratamento',
    'Título da Disponibilidade': 'titulo_disponibilidade', 'Largura da Disponibilidade': 'largura_disponibilidade',
    'Legenda de destaque': 'legenda_destaque', 'Selo de montagem': 'selo_montagem', 'Bloco compacto': 'compacta',
    'Código': 'cod', 'Índice': 'indice', 'Nome do produto': 'produto', 'Blue UV': 'blue_uv', 'Ponto': 'ponto',
    'Destaque': 'destaque', 'Tratamento': 'tratamento', 'Pontos do tratamento': 'pontos_tratamento',
    'Esf. máx': 'esf_max', 'Esf. mín': 'esf_min', 'Cil.': 'cil', 'Add.': 'add', 'Diâm.': 'diam', 'Curva': 'curva',
    'Linha extra': 'disp_extra', 'Preço 1': 'preco_1', 'Preço 2': 'preco_2', 'Preço 3': 'preco_3', 'Preço 4': 'preco_4',
    'Obs. do preço': 'obs_preco', 'Cores e códigos': 'cores', 'Página': 'pagina', 'Bloco': 'bloco', 'Valor': 'valor',
    'Ícone': 'icone', 'Marca': 'marca', 'Produto': 'produto', 'Montado': 'montado', 'Par + montagem': 'montado',
    'Bolinha': 'bolinha', 'Observação': 'obs',
}
ABAS = {'Config': ('Mês', 'Config'), 'Paginas': ('Páginas', 'Paginas'), 'Familias': ('Famílias', 'Familias'), 'Produtos': ('Produtos',), 'Combos': ('Combos',)}
PRIMEIRAS = {'chave', 'Campo', 'familia', 'Família', 'pagina', 'Página', 'Código', 'cod'}

def aba(wb, nome, obrigatoria=True):
    for n in ABAS[nome]:
        if n in wb.sheetnames: return wb[n]
    if not obrigatoria: return None
    sys.exit(f'ERRO: aba "{ABAS[nome][0]}" não existe no Excel.')

def tabela(ws):
    linhas = list(ws.iter_rows(values_only=True))
    # O título pode estar na linha 1 ou 2 (no modelo novo a linha 1 é a faixa de grupos).
    h = next(i for i, l in enumerate(linhas[:5]) if l and txt(l[0]) in PRIMEIRAS)
    cab = [ALIAS.get(txt(c), txt(c)) for c in linhas[h]]
    out = []
    for i, l in enumerate(linhas[h + 1:], start=h + 2):
        d = {cab[j]: l[j] for j in range(len(cab)) if cab[j]}
        if all(v is None or txt(v) == '' for v in d.values()): continue
        d['_linha'] = i
        out.append(d)
    return out

def cores(v, onde):
    """'8110 Cinza; 81111 Marrom' -> [['8110','Cinza'],...]; 'G15' -> [['','G15']]"""
    if not txt(v): return []
    res = []
    for parte in re.split(r'[;|]', txt(v)):
        parte = parte.strip()
        if not parte: continue
        m = re.match(r'^(\d+)\s+(.+)$', parte)
        cod, nome = (m.group(1), m.group(2).strip()) if m else ('', parte)
        if nome not in CONHECIDAS: avisos.append(f'{onde}: cor "{nome}" fora da paleta — vai sair cinza.')
        res.append([cod, nome])
    nomes = [n for _, n in res]
    for seq in (GEN_S, COLORS, ESPELHADO):
        pos = [seq.index(n) for n in nomes if n in seq]
        if len(pos) > 1 and pos != sorted(pos):
            erros.append(f'{onde}: cores fora da ordem fixa ({", ".join(nomes)}). Ordem: {" · ".join(seq)}.')
    return res

def icone_de(F):
    """Mesma regra do construtor.js: coluna Ícone do Excel ou, vazio, o nome da família."""
    e = (F.get('icone') or '').lower()
    if e == 'nenhum': return None
    if e: return e
    n = F['familia'].upper()
    if 'FREEVIX' in n: return 'freevix'
    if 'ESSILOR' in n: return 'essilor'
    if 'OPTVIEW' in n: return 'sol'
    if 'ESPACE' in n: return 'brilho'
    if 'PRONTA' in n: return 'caixa'
    if 'KODAK' in n: return 'olho'
    if re.search(r'VIX TOTAL|OPTF|MULTIFOCAL', n): return 'oculos'
    return None

def ar_de(titulo):
    k = titulo.lower()
    if 'blue' in k: return 'blue'
    if 'guard' in k: return 'guard'
    if 'express' in k: return 'express'
    return None

def svg_qr(url, aqui):
    """SVG do QR (retângulos cheios, que o Figma importa sem erro de traço). Usa o pacote segno quando existe;
    sem ele, só serve o endereço padrão, que vai pronto em modelo/svg/qr-segundo-par-figma.svg."""
    try:
        import segno
        m = [list(r) for r in segno.make(url, error='q', boost_error=False).matrix]
        n, d = len(m), []
        for y, row in enumerate(m):
            x = 0
            while x < n:
                if row[x]:
                    w = 1
                    while x + w < n and row[x + w]: w += 1
                    d.append('M%d %dh%dv1h-%dz' % (x, y, w, w)); x += w
                else: x += 1
        return '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 %d %d" shape-rendering="crispEdges"><path fill="#000000" d="%s"/></svg>' % (n, n, ''.join(d))
    except ImportError:
        if url == QR_URL_PADRAO:
            f = aqui / 'modelo' / 'svg' / 'qr-segundo-par-figma.svg'
            if f.exists(): return f.read_text(encoding='utf-8').strip()
        return None

def main():
    try: sys.stdout.reconfigure(encoding="utf-8")
    except Exception: pass
    ap = argparse.ArgumentParser()
    ap.add_argument('excel'); ap.add_argument('--saida', default='paginas')
    a = ap.parse_args()
    wb = load_workbook(a.excel, data_only=True)
    abas = {n: aba(wb, n) for n in ABAS if n != 'Combos'}
    aba_combos = aba(wb, 'Combos', False)

    cfg = {txt(r['chave']): r.get('valor') for r in tabela(abas['Config']) if txt(r.get('chave'))}
    for k in ('mes', 'ano', 'validade'):
        if not txt(cfg.get(k)): erros.append(f'Config: "{k}" vazio.')
    ag, ab = num(cfg.get('acrescimo_guard')), num(cfg.get('acrescimo_blue'))

    fams = {}
    for r in tabela(abas['Familias']):
        nome = txt(r.get('familia'))
        if not nome or not txt(r.get('cor')): continue          # linhas de ajuda
        titulos = [t.strip() for t in txt(r.get('colunas_preco')).split(';') if t.strip()]
        if not 1 <= len(titulos) <= 4: erros.append(f'Familias l.{r["_linha"]} {nome}: colunas_preco precisa de 1 a 4 títulos.')
        cor = txt(r.get('cor')).upper()
        if not re.match(r'^#[0-9A-F]{6}$', cor): erros.append(f'Familias l.{r["_linha"]} {nome}: cor "{cor}" não é hex (#RRGGBB).')
        # Título curto numa linha só; longo quebra no primeiro espaço (Reflecta / Guard).
        partes = lambda t: t.split(' ', 1) if len(t) > 9 and ' ' in t else [t, '']
        fams[nome] = dict(
            familia=nome, subtitulo=txt(r.get('tipo')), cor=cor,
            pilulas=[p.strip() for p in txt(r.get('pilulas')).split(';') if p.strip()],
            titulos=[partes(t) for t in titulos],
            titulosTexto=titulos,
            promo=[int(x) for x in re.findall(r'\d', txt(r.get('colunas_promo')))],
            regra50=sn(r.get('regra_50')), sep=sn(r.get('separador_indice')), trat=sn(r.get('coluna_tratamento')),
            dispTitulo=txt(r.get('titulo_disponibilidade')) or None,
            dispW=int(num(r.get('largura_disponibilidade')) or 70),
            legenda=sn(r.get('legenda_destaque')), montagem=sn(r.get('selo_montagem')),
            compacta=sn(r.get('compacta')), icone=txt(r.get('icone')).lower(), rows=[])
        if fams[nome]['compacta'] and len(titulos) > 1:
            erros.append(f'Familias l.{r["_linha"]} {nome}: compacta = S só funciona com 1 coluna de preço.')

    vistos = {}
    for r in tabela(abas['Produtos']):
        fam = txt(r.get('familia')); onde = f'Produtos l.{r["_linha"]}'
        if fam not in fams: erros.append(f'{onde}: família "{fam}" não está na aba Familias.'); continue
        F = fams[fam]
        cod, ind, nome = txt(r.get('cod')), txt(r.get('indice')), txt(r.get('produto'))
        # Excel costuma virar 0357 em 357 e 1.50 em 1.5 quando a célula não está como texto.
        if isinstance(r.get('cod'), (int, float)) and len(cod) < 4: cod = cod.zfill(4)
        if isinstance(r.get('indice'), (int, float)): ind = f'{float(r["indice"]):.2f}'
        if not nome: erros.append(f'{onde}: produto vazio.')
        if not re.match(r'^\d\.\d\d$', ind): erros.append(f'{onde}: índice "{ind}" — use 1.50, 1.59, 1.67…')
        cs = cores(r.get('cores'), onde)
        if not cod and len(cs) < 2: erros.append(f'{onde}: cod vazio sem linha de cores (2+ códigos) — a seta ↓ não teria para onde apontar.')
        for c in [cod] + [c for c, _ in cs]:
            if not c: continue
            if c in vistos and vistos[c] != onde: avisos.append(f'{onde}: código {c} repetido (já em {vistos[c]}).')
            vistos.setdefault(c, onde)
        # preços
        n = len(F['titulos'])
        ps = [num(r.get(f'preco_{i}')) for i in range(1, n + 1)]
        for i in range(n + 1, 5):
            if num(r.get(f'preco_{i}')) is not None: avisos.append(f'{onde}: preco_{i} preenchido mas a família só tem {n} colunas — ignorado.')
        if F['regra50']:
            if ag is None or ab is None: erros.append('Config: acrescimo_guard / acrescimo_blue vazios, e há família com regra_50.')
            elif ps[0] is not None and n >= 3:
                esperado = {1: ps[0] + ag, 2: ps[0] + ab}
                for i, v in esperado.items():
                    if ps[i] is None: ps[i] = round(v, 2)
                    elif abs(ps[i] - v) > 0.005:
                        erros.append(f'{onde} {cod or nome}: {F["titulosTexto"][i]} = {reais(ps[i])}, a regra 50% dá {reais(v)}.')
            elif ps[0] is None and n >= 3 and (ps[1] is None or ps[2] is None):
                erros.append(f'{onde} {cod or nome}: sem Sem A.R., Guard e Blue precisam vir preenchidos.')
        vals = [v for v in ps if v is not None]
        if vals != sorted(vals): avisos.append(f'{onde} {cod or nome}: preços fora de ordem crescente ({", ".join(reais(v) for v in vals)}).')
        if not vals: erros.append(f'{onde}: linha sem nenhum preço.')
        obs = {}
        m = re.match(r'^\s*(\d)\s*:\s*(.+)$', txt(r.get('obs_preco')))
        if m: obs[int(m.group(1))] = m.group(2).strip()
        elif txt(r.get('obs_preco')): erros.append(f'{onde}: obs_preco no formato "2: texto".')
        # disponibilidade
        emax, emin = num(r.get('esf_max')), num(r.get('esf_min'))
        curva, diam = txt(r.get('curva')), txt(r.get('diam'))
        disp = []
        if curva: disp.append(f'Curva {curva}' + (f' | Ø{diam}' if diam else ''))
        elif emax is not None and emin is not None:
            if emax < emin: emax, emin = emin, emax
            disp.append(f'Esf. {dec(emax)} a {dec(emin)}')
        l2 = []
        if txt(r.get('cil')): l2.append('Cil. ' + txt(r.get('cil')))
        if txt(r.get('add')): l2.append('Add. ' + txt(r.get('add')))
        if diam and not curva: l2.append('Ø' + diam)
        if l2: disp.append(' | '.join(l2))
        if txt(r.get('disp_extra')): disp.append(txt(r.get('disp_extra')))
        if not disp: avisos.append(f'{onde}: disponibilidade vazia.')
        dest = txt(r.get('destaque')).upper() or None
        if dest and dest not in ('FOTO', 'RX'): erros.append(f'{onde}: destaque "{dest}" — use FOTO ou RX.')
        ponto = {'verde': 'verde', 'azul': 'azul', 'roxo': 'roxo'}.get(txt(r.get('ponto')).lower()) if txt(r.get('ponto')) else None
        F['rows'].append(dict(
            c=cod, i=ind, n=nome, b=sn(r.get('blue_uv')), d=ponto,
            t=txt(r.get('tratamento')) or None,
            td=[p.strip().lower() for p in txt(r.get('pontos_tratamento')).split(';') if p.strip()],
            disp=disp, p=[reais(v) if v is not None else '' for v in ps], obs=obs, dest=dest, cores=cs, _preco=ps[0]))

    # combos (opcional): aba Combos + chaves combo_* no Mês + bloco COMBO em Páginas
    combo_ativo = sn(cfg.get('combo_ativo'))
    combos = []
    precos_tabela = {}                       # código -> menor preço da 1ª coluna nas tabelas normais
    for F in fams.values():
        for r in F['rows']:
            if r['c'] and r['_preco'] is not None: precos_tabela[r['c']] = min(r['_preco'], precos_tabela.get(r['c'], 1e12))
    if aba_combos is not None:
        for r in tabela(aba_combos):
            onde = f'Combos l.{r["_linha"]}'
            cod, marca, nome = txt(r.get('cod')), txt(r.get('marca')), txt(r.get('produto'))
            if isinstance(r.get('cod'), (int, float)) and len(cod) < 4: cod = cod.zfill(4)
            if not cod and not nome: continue
            if not cod or not nome: erros.append(f'{onde}: código e produto são obrigatórios.'); continue
            if marca not in MARCAS_COMBO: erros.append(f'{onde} {cod}: marca "{marca}" — use {", ".join(sorted(MARCAS_COMBO))}.')
            try: preco = num(r.get('montado'))
            except Exception: preco = None
            if preco is None: erros.append(f'{onde} {cod}: preço montado vazio ou inválido.'); continue
            bol = txt(r.get('bolinha')).lower()
            if bol and bol not in ('verde', 'azul', 'roxo'): erros.append(f'{onde} {cod}: bolinha "{bol}" — use verde, azul ou roxo.'); bol = ''
            if marca and nome.lower().startswith(marca.lower() + ' '): avisos.append(f'{onde} {cod}: o produto começa com a marca ("{marca}"); ela já vira chip, tire do nome.')
            if cod in precos_tabela and preco < precos_tabela[cod] - 0.005:
                avisos.append(f'{onde} {cod}: combo montado ({reais(preco)}) abaixo do preço da lente na tabela ({reais(precos_tabela[cod])}). A montagem entra no combo, confira.')
            if cod not in precos_tabela: avisos.append(f'{onde} {cod}: código não está em nenhuma tabela normal (ok se for só do combo).')
            combos.append(dict(c=cod, marca=marca, nome=nome, preco=reais(preco), bol=bol))
    codigos_combo = {c['c'] for c in combos}
    for F in fams.values():
        for r in F['rows']:
            if r['c'] in codigos_combo: r['k'] = True
    painel = txt(cfg.get('combo_painel')).upper() or 'NENHUM'
    if painel not in PAINEIS: erros.append(f'Mês: combo_painel "{painel}" — use {", ".join(sorted(PAINEIS))}.')

    # páginas
    pags = {}
    for r in tabela(abas['Paginas']):
        b = txt(r.get('bloco'))
        if not b: continue
        try: n = int(num(r.get('pagina')))
        except Exception: erros.append(f'Paginas l.{r["_linha"]}: página "{txt(r.get("pagina"))}" inválida.'); continue
        if b not in ESPECIAIS and b not in fams: erros.append(f'Paginas l.{r["_linha"]}: bloco "{b}" não é família nem bloco especial.')
        pags.setdefault(n, []).append(b)
    usadas = {b for bs in pags.values() for b in bs}
    if 'COMBO' in usadas and not combo_ativo: erros.append('Páginas tem o bloco COMBO, mas Mês: combo_ativo não está em S.')
    if combo_ativo and 'COMBO' not in usadas: erros.append('Mês: combo_ativo = S, mas o bloco COMBO não está em nenhuma página da aba Páginas.')
    if 'COMBO' in usadas and not combos: erros.append('Bloco COMBO em Páginas, mas a aba Combos está vazia.')
    if 'COMBO' in usadas and painel == 'VIXCLUB' and 'VIXCLUB' in usadas: erros.append('Painel VIXCLUB do combo e bloco VIXCLUB avulso na mesma promoção: use só um.')
    for f in fams:
        if f not in usadas and fams[f]['rows']: erros.append(f'Família "{f}" tem produtos mas não está em nenhuma página.')
    for f, F in fams.items():
        if f in usadas and not F['rows']: erros.append(f'Família "{f}" está na aba Paginas mas não tem produtos.')
    numeros = sorted(pags)
    if numeros != list(range(1, len(numeros) + 1)): erros.append(f'Paginas: numeração com buraco ({numeros}).')
    if len(numeros) % 4: avisos.append(f'{len(numeros)} páginas — não é múltiplo de 4 (impressão).')

    # estimativa de altura (a verdade é o fimDoConteudo que o construtor devolve)
    def altura_familia(F):
        if F['compacta']: return 13 + 22 + 4 + 20 * ((len(F['rows']) + 1) // 2)
        h = 58 + 8 + 18 + 22 + 3          # cabeçalho da família, gap, padding da tabela, títulos, régua
        if F['promo'] and txt(cfg.get('adesivo_promo')): h += 10     # folga sob o adesivo −50%
        ind = None
        for r in F['rows']:
            if F['sep'] and r['i'] != ind: h += 25
            ind = r['i']
            h += max(14, 7.5 * len(r['disp']) + 1) + 6
            if len(r['cores']) > 1: h += 18
        return h
    estim = {}
    for n in numeros:
        y = TOPO_CAPA if 'CAPA' in pags[n] else TOPO
        blocos = [b for b in pags[n] if b != 'CAPA']
        for i, b in enumerate(blocos):
            if i: y += 14
            y += {'FAIXA_50': 53, 'FAIXA_MONTAGEM': 53, 'VIXCLUB': 150, 'COMBO': 640}.get(b) or altura_familia(fams[b])
        estim[n] = round(y)
        # O construtor aperta até ~30px sozinho; acima disso, a página não cabe.
        if y > LIMITE + 30: avisos.append(f'Página {n}: estimativa {round(y)} bem acima de {LIMITE} — provavelmente estoura mesmo com o aperto. Considere mover um bloco.')

    # saída
    fonte = (AQUI / 'construtor.js').read_text(encoding='utf-8')
    svg_vix = (AQUI / 'modelo' / 'svg' / 'vixlens-negativo.svg').read_text(encoding='utf-8')
    svg_club = (AQUI / 'modelo' / 'svg' / 'vixclub.svg').read_text(encoding='utf-8')
    svg_oferta = (AQUI / 'modelo' / 'svg' / 'selo-oferta.svg').read_text(encoding='utf-8').strip()
    icones = json.loads((AQUI / 'modelo' / 'svg' / 'icones.json').read_text(encoding='utf-8'))
    qr_url = txt(cfg.get('combo_qr_url')) or QR_URL_PADRAO
    qr_svg = None
    if 'COMBO' in usadas and painel == 'CAMPANHA_FREEVIX':
        qr_svg = svg_qr(qr_url, AQUI)
        if qr_svg is None: erros.append(f'QR: o endereço {qr_url} é diferente do padrão e o pacote segno não está instalado (pip install segno).')
    mes, ano = txt(cfg.get('mes')).upper(), txt(cfg.get('ano'))
    contato = '   ·   '.join(x for x in (txt(cfg.get('instagram')), txt(cfg.get('telefone')), txt(cfg.get('site'))) if x)
    saida = Path(a.saida); saida.mkdir(parents=True, exist_ok=True)
    total_prod = sum(len(F['rows']) for F in fams.values())
    total_cor = sum(sum(1 for c, _ in r['cores'] if c) for F in fams.values() for r in F['rows'])
    resumo = dict(mes=mes, ano=ano, paginas=len(numeros), produtos=total_prod, codigosCor=total_cor,
                  estimativa=estim, porPagina={}, erros=erros, avisos=avisos)
    paginas_js = []
    for n in numeros:
        blocos = []
        for b in pags[n]:
            if b == 'CAPA': blocos.append({'tipo': 'capa'})
            elif b == 'FAIXA_50': blocos.append({'tipo': 'faixa50', 'antes': txt(cfg.get('faixa50_antes')), 'destaque': txt(cfg.get('faixa50_destaque')), 'depois': txt(cfg.get('faixa50_depois')), 'direita': txt(cfg.get('faixa50_direita'))})
            elif b == 'FAIXA_MONTAGEM': blocos.append({'tipo': 'montagem', 'antes': txt(cfg.get('montagem_antes')), 'destaque': txt(cfg.get('montagem_destaque')), 'depois': '', 'direita': txt(cfg.get('montagem_direita'))})
            elif b == 'VIXCLUB': blocos.append({'tipo': 'vixclub', 'titulo': txt(cfg.get('vixclub_titulo')), 'texto': txt(cfg.get('vixclub_texto'))})
            elif b == 'COMBO':
                blocos.append({'tipo': 'combo', 'titulo': txt(cfg.get('combo_titulo')), 'chamada': txt(cfg.get('combo_chamada')), 'chip': txt(cfg.get('combo_chip')),
                               'colunaPreco': txt(cfg.get('combo_coluna_preco')) or 'Montado', 'itens': combos, 'painel': painel,
                               'chamadaCampanha': txt(cfg.get('combo_chamada_campanha')) or 'Fale com seu consultor e aponte a câmera para o QR Code para conferir as regras e saber mais.',
                               'qrSvg': qr_svg, 'qrUrl': qr_url})
                if painel == 'VIXCLUB': blocos.append({'tipo': 'vixclub', 'titulo': txt(cfg.get('vixclub_titulo')), 'texto': txt(cfg.get('vixclub_texto'))})
            else:
                F = dict(fams[b]); F.pop('regra50')
                F['rows'] = [{k: v for k, v in r.items() if k != '_preco'} for r in F['rows']]
                blocos.append({'tipo': 'familia', **F})
        pag = dict(numero=n, total=len(numeros), mes=mes, ano=ano, validade=txt(cfg.get('validade')).upper(),
                   aviso=txt(cfg.get('aviso')), contato=contato, adesivo=txt(cfg.get('adesivo_promo')),
                   comboSelo=txt(cfg.get('combo_selo')) or 'Combo disponível',
                   montagem=dict(rotulo=txt(cfg.get('montagem_selo')) or 'MONTAGEM\nLENTE PRONTA', valor=txt(cfg.get('montagem_destaque')))
                   if any(fams[b]['montagem'] for b in pags[n] if b in fams) else None,
                   frameNome=f'Promovix {mes} // P{n:02d}', x=(n - 1) * 640, blocos=blocos,
                   svgVixlens=svg_vix if 'CAPA' in pags[n] else None,
                   svgVixclub=svg_club if ('VIXCLUB' in pags[n] or ('COMBO' in pags[n] and painel == 'VIXCLUB')) else None)
        ph, ar, marca = {}, {}, {}
        for b in pags[n]:
            if b == 'COMBO': ph['oculos'] = icones['ph']['oculos']
            if b not in fams or fams[b]['compacta']: continue
            F = fams[b]; k = icone_de(F)
            if k in icones['marca']: marca[k] = icones['marca'][k]
            elif k in icones['ph']: ph[k] = icones['ph'][k]
            elif k: avisos.append(f'Famílias {b}: ícone "{k}" desconhecido (use oculos, olho, sol, caixa, brilho, freevix, essilor ou nenhum).')
            for t in F['titulosTexto']:
                a = ar_de(t)
                if a: ar[a] = icones['ar'][a]
            if any(r.get('k') for r in F['rows']): ph['selo'] = icones['ph']['selo']
        pag.update(ph=ph, ar=ar, marca=marca)
        paginas_js.append((n, pag))
        resumo['porPagina'][n] = dict(blocos=pags[n], produtos=sum(len(fams[b]['rows']) for b in pags[n] if b in fams))

    # Agrupa páginas em lotes: o construtor vai uma vez por lote, e o lote fica
    # abaixo do limite de 50 mil caracteres do use_figma.
    def secoes(lista):
        tipos, compacta, montagem = set(), False, False
        for _, pg in lista:
            for b in pg['blocos']:
                tipos.add(b['tipo'])
                if b['tipo'] == 'familia' and b.get('compacta'): compacta = True
            if pg.get('montagem'): montagem = True
        return {'capa': 'capa' in tipos, 'combo': 'combo' in tipos, 'vixclub': 'vixclub' in tipos,
                'faixa': bool(tipos & {'faixa50', 'montagem'}), 'compacta': compacta, 'montagem': montagem}
    def empacotar(lista):
        js = fonte
        for nome, usa in secoes(lista).items():
            if not usa: js = re.sub(r'//#sec ' + nome + r'\n.*?//#fim ' + nome + r'\n', '', js, flags=re.S)
        js = js.replace('/*PAGINAS*/[]', json.dumps([p for _, p in lista], ensure_ascii=False, separators=(',', ':')))
        js = js.replace('/*SVG_OFERTA*/null', json.dumps(svg_oferta))
        return '\n'.join(l.strip() for l in js.splitlines() if l.strip() and not l.strip().startswith('//'))
    for f in saida.glob('lote_*.js'): f.unlink()
    lotes, atual = [], []
    for item in paginas_js:
        if atual and len(empacotar(atual + [item])) > 46000:
            lotes.append(atual); atual = []
        atual.append(item)
    if atual: lotes.append(atual)
    resumo['lotes'] = []
    for k, lote in enumerate(lotes, start=1):
        js = empacotar(lote)
        if len(js) > 49000: erros.append(f'Lote {k}: {len(js)} caracteres, acima do limite do use_figma.')
        (saida / f'lote_{k:02d}.js').write_text(js, encoding='utf-8')
        resumo['lotes'].append(dict(arquivo=f'lote_{k:02d}.js', paginas=[n for n, _ in lote], caracteres=len(js)))
    (saida / 'resumo.json').write_text(json.dumps(resumo, ensure_ascii=False, indent=1), encoding='utf-8')

    print(f'Promovix {mes}.{ano}: {len(numeros)} páginas, {total_prod} produtos, {total_cor} códigos de cor.')
    for n in numeros: print(f'  P{n:02d} estimativa {estim[n]:>4}  {" + ".join(pags[n])}')
    for l in resumo['lotes']: print(f'  {l["arquivo"]}: páginas {l["paginas"]} ({l["caracteres"]} caracteres)')
    for e in erros: print('ERRO  ', e)
    for w in avisos: print('AVISO ', w)
    sys.exit(1 if erros else 0)

if __name__ == '__main__':
    main()
